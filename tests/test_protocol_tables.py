"""Детерминированный извлекатель таблиц состава из протоколов
(parsers/protocol_tables) — замечание эколога 08.09.2026: состав в паспорте
без «Прочих компонентов», неполное распознавание помечается, таблица
главнее ИИ-строк, сопоставление протокол ↔ отход по ФККО/наименованию."""
from pathlib import Path

import fitz
import pytest
from docx import Document

from ecodoc.core import waste_refdata as R
from ecodoc.core.models import ReportContext, WasteFlow
from ecodoc.parsers import protocol_tables as PT

HEAD = ["№ п/п", "Наименование определяемого показателя (компонента)",
        "Единицы измерения", "Результат измерения ± значение неопределенности",
        "Наименование (шифр) методики измерения (МИ)"]
TBO = ("Мусор от офисных и бытовых помещений организаций несортированный "
       "(исключая крупногабаритный)")


def _pdf_with_table(path: Path, rows: list[list[str]], head_lines: list[str],
                    second_page: list[list[str]] | None = None) -> Path:
    """PDF с текстовым слоем и таблицей, нарисованной линиями (find_tables)."""
    doc = fitz.open()
    font = fitz.Font("cjk")

    def _page(lines, table_rows):
        page = doc.new_page()
        tw = fitz.TextWriter(page.rect)
        y = 40
        for line in lines:
            tw.append((40, y), line, font=font, fontsize=8)
            y += 12
        y0 = y + 10
        xs = [40, 70, 300, 340, 430, 560]
        ys = [y0 + 16 * i for i in range(len(table_rows) + 1)]
        for x in xs:
            page.draw_line((x, ys[0]), (x, ys[-1]))
        for yy in ys:
            page.draw_line((xs[0], yy), (xs[-1], yy))
        for i, r in enumerate(table_rows):
            for j, c in enumerate(r):
                tw.append((xs[j] + 2, ys[i] + 11), c, font=font, fontsize=6)
        tw.write_text(page)

    _page(head_lines, [HEAD] + rows)
    if second_page is not None:
        _page(["Продолжение таблицы"], [HEAD] + second_page)
    doc.save(str(path))
    return path


HEAD_LINES = [
    "ИЦ ООО «ТАСИС»  Уникальный номер записи об аккредитации: № РОСС RU.0001.21АУ50",
    "ПРОТОКОЛ ИССЛЕДОВАНИЙ (ИЗМЕРЕНИЙ) № 20002.25-1-Отх от 27 февраля 2025 г.",
    "Объект исследований (измерений)*: Отходы",
    "Наименование вида отхода*: " + TBO,
    "Код по ФККО: 7 33 100 01 72 4",
    "Результаты исследований (измерений)",
]


def test_pdf_table_full_composition(tmp_path):
    """Таблица PDF (find_tables) → запись состава с реквизитами; сумма 100 —
    полный состав, note пустой."""
    rows = [["1", "Бумага", "%", "26,1 ± 0,3", "М-27-2023"],
            ["2", "Картон", "%", "60,9 ± 0,5", "М-27-2023"],
            ["3", "Полиэтилен", "%", "13,0 ± 0,2", "М-27-2023"]]
    pdf = _pdf_with_table(tmp_path / "прот.ТБО.pdf", rows, HEAD_LINES)
    recs = PT.extract(pdf)
    assert len(recs) == 1
    r = recs[0]
    assert r["method"] == "table" and r["kind"] == PT.KIND
    assert r["protocol_no"] == "20002.25-1-Отх" and r["date"] == "27.02.2025"
    assert r["fkko"] == "73310001724"
    assert r["object"].startswith("Мусор от офисных и бытовых помещений")
    assert r["lab"] == "ИЦ ООО «ТАСИС»" and "21АУ50" in r["lab_attestation"]
    assert r["method_doc"] == "М-27-2023"
    assert [(s["name"], s["value"], s["unit"]) for s in r["substances"]] == [
        ("Бумага", 26.1, "%"), ("Картон", 60.9, "%"), ("Полиэтилен", 13.0, "%")]
    assert r["total_pct"] == 100.0 and r["incomplete"] is False and r["note"] == ""
    assert r["_src"] == "прот.ТБО.pdf (лист 1)" and r["page"] == 1


def test_pdf_table_continues_on_next_page(tmp_path):
    """Продолжение таблицы на следующем листе (без своей шапки протокола)
    дописывается к той же записи."""
    rows = [["1", "Бумага", "%", "60,0", "М-27-2023"]]
    more = [["2", "Картон", "%", "40,0", "М-27-2023"]]
    pdf = _pdf_with_table(tmp_path / "п.pdf", rows, HEAD_LINES, second_page=more)
    recs = PT.extract(pdf)
    assert len(recs) == 1
    assert [s["name"] for s in recs[0]["substances"]] == ["Бумага", "Картон"]
    assert recs[0]["total_pct"] == 100.0 and "листы 1–2" in recs[0]["_src"]


def test_docx_table_incomplete_no_others(tmp_path):
    """Таблица .docx: сумма 61,4 % → incomplete + причина (сколько строк,
    какой лист), «Прочих» ни в записи, ни после normalize_components."""
    d = Document()
    d.add_paragraph("ПРОТОКОЛ ИССЛЕДОВАНИЙ (ИЗМЕРЕНИЙ) № 13208.26-1-Отх от 17.08.2026")
    d.add_paragraph("Наименование вида отхода: " + TBO)
    d.add_paragraph("Код по ФККО: 7 33 100 01 72 4")
    rows = [["1", "Картон", "мг/кг", "169000"], ["2", "Пищевые отходы", "мг/кг", "102000"],
            ["3", "Полипропилен", "мг/кг", "101000"],
            ["4", "Лом цветных металлов (алюминий)", "мг/кг", "81000"],
            ["5", "Текстиль (х/б)", "мг/кг", "69000"], ["6", "Полиэтилен", "мг/кг", "67000"],
            ["7", "Прочее (неклассифицируемые материалы)", "мг/кг", "25000"]]
    t = d.add_table(rows=len(rows) + 1, cols=4)
    for j, h in enumerate(HEAD[:4]):
        t.cell(0, j).text = h
    for i, r in enumerate(rows, 1):
        for j, c in enumerate(r):
            t.cell(i, j).text = c
    path = tmp_path / "протокол.docx"
    d.save(path)
    recs = PT.extract(path)
    assert len(recs) == 1
    r = recs[0]
    assert r["protocol_no"] == "13208.26-1-Отх" and r["date"] == "17.08.2026"
    assert r["fkko"] == "73310001724"
    assert len(r["substances"]) == 7 and r["substances"][0]["unit"] == "мг/кг"
    assert r["total_pct"] == 61.4 and r["incomplete"] is True
    assert "распознано 61.4 % (7 строк)" in r["note"]
    assert "№ 13208.26-1-Отх" in r["note"] and "протокол.docx, лист 1" in r["note"]
    assert not any("неидентифицированные" in s["name"] for s in r["substances"])
    # «Прочее (неклассифицируемые материалы)» — строка САМОГО протокола, остаётся
    comps, note = R.normalize_components(PT.as_components(r), source="протокол")
    assert [c["name"] for c in comps][-1] == "Прочее (неклассифицируемые материалы)"
    assert R.is_incomplete_note(note) and "61.4 %" in note
    assert not any(R.is_synthetic_row(c["name"]) for c in comps)
    assert R.components_total(comps) == 61.4          # как есть, не «до 100»


def test_text_pages_rows_and_unnamed(tmp_path):
    """Кэш текстов (без исходника): строки регэкспом; строка с величиной без
    наименования (OCR) считается и попадает в note."""
    page = "\n".join([
        "ПРОТОКОЛ ИССЛЕДОВАНИЙ (ИЗМЕРЕНИЙ) № 26212.25-1-Отх от 19 декабря 2025 г.",
        "Наименование вида отхода*: Осадок (шлам) механической очистки "
        "нефтесодержащих сточных вод, содержащий нефтепродукты в количестве менее 15 %",
        "Результаты исследований (измерений)",
        "Наименование определяемого показателя (компонента) Результат измерения",
        "1 Массовая доля нефтепродуктов % 11,5 ± 3,7 ПНД Ф 16.1:2:2.2:2.3:3.64-10",
        "28,9 ± 2,0 ПНД Ф 16.1:2.2:2.3:3.58-08",
        "3 Массовая доля золы (зольности) % 59,6 ± 2,0 ПНД Ф 16.2.2:2.3:3.29-02",
        "Примечания:",
        "1. Отклонения от указанных МИ не установлены.",
    ])
    recs = PT.from_text_pages([page], "26212_25.pdf")
    assert len(recs) == 1
    r = recs[0]
    assert [(s["name"], s["value"]) for s in r["substances"]] == [
        ("Массовая доля нефтепродуктов", 11.5), ("Массовая доля золы (зольности)", 59.6)]
    assert r["rows_unnamed"] == 1 and r["incomplete"]
    assert "2 строки, ещё 1 строка с величиной без наименования" in r["note"]
    assert r["object"].startswith("Осадок (шлам)")


def test_xlsx_table(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Протокол № 5-Отх от 01.03.2026"])
    ws.append(["Наименование вида отхода: Лом бетонных изделий, отходы бетона в кусковой форме"])
    ws.append(["Код ФККО", "8 22 201 01 21 5"])
    ws.append(["№", "Компонент", "Содержание, %"])
    ws.append([1, "Бетон", 100])
    path = tmp_path / "п.xlsx"
    wb.save(path)
    (r,) = PT.extract(path)
    assert r["fkko"] == "82220101215" and r["protocol_no"] == "5-Отх"
    assert r["substances"] == [{"name": "Бетон", "value": 100.0, "unit": "%"}]
    assert not r["incomplete"]


def test_names_match_and_match_waste():
    """Сопоставление по наименованию: регистр, ё, пунктуация, первые 40
    символов либо ≥ 70 % общих слов длиннее 3 букв; по ФККО — точное."""
    assert PT.names_match(TBO, "мусор от офисных и бытовых помещений организаций "
                               "несортированный, исключая крупногабаритный")
    assert PT.names_match("Лом чёрных металлов несортированный",
                          "лом черных металлов, несортированный")
    assert not PT.names_match("Мусор от офисных и бытовых помещений",
                              "Отходы (мусор) от строительных и ремонтных работ")
    rec = {"fkko": "73310001724", "object": TBO}
    assert PT.match_waste(rec, "7 33 100 01 72 4", "")
    assert not PT.match_waste(rec, "8 90 000 01 72 4", TBO)      # свой код — другой
    rec2 = {"fkko": "", "object": TBO}
    assert PT.match_waste(rec2, "7 33 100 01 72 4", "Мусор от офисных и бытовых "
                                                    "помещений организаций несортированный")
    assert not PT.match_waste(rec2, "", "Осадок мойки")
    # protocol_matches_waste (отбор протокола для паспорта) — те же правила
    assert R.protocol_matches_waste({"kind": "состав/КХА", "object": TBO,
                                     "substances": [{"unit": "%"}]},
                                    "", "мусор от офисных и бытовых помещений "
                                        "организаций несортированный")


def test_merge_into_table_beats_ai_rows():
    """Табличная запись заменяет ИИ-строки того же протокола/файла, забирая
    у них дату/лабораторию/аттестат; повторный разбор заменяет прежнюю."""
    ctx = ReportContext()
    ctx.extra["lab_results"] = [
        {"kind": "КХА", "protocol_no": "13208.26-1-Отх", "date": "17.08.2026",
         "lab": "ИЦ ООО «ТАСИС»", "lab_attestation": "№ РОСС RU.0001.21АУ50",
         "object": TBO, "method": "М-27-2023",
         "substances": [{"name": "Картон", "value": "169000", "unit": "мг/кг"}],
         "_src": "13208_26.pdf (листы 1–3)"},
        {"kind": "биотест", "protocol_no": "13308.26-1-ТОтх", "object": TBO,
         "_src": "13208_26.pdf (листы 37–41)"},
        {"kind": "КХА", "protocol_no": "13208.26-1-Отх", "object": TBO,
         "substances": [], "_src": "другой.pdf (лист 1)"}]
    rec = {"kind": PT.KIND, "protocol_no": "13208 .26 -1 -Отх", "date": "", "lab": "TACHC",
           "lab_attestation": "", "object": TBO, "fkko": "73310001724",
           "method": "table", "method_doc": "",
           "substances": [{"name": "Бумага", "value": 26.1, "unit": "%"}],
           "_src": "13208_26.pdf (лист 2)", "page": 2, "total_pct": 26.1,
           "incomplete": True, "note": ""}
    res = PT.merge_into(ctx, [rec])
    assert res == {"added": 1, "replaced": 1}
    labs = ctx.extra["lab_results"]
    assert len(labs) == 3
    t = next(l for l in labs if l.get("method") == "table")
    assert t["date"] == "17.08.2026" and t["lab"] == "ИЦ ООО «ТАСИС»"
    assert "21АУ50" in t["lab_attestation"] and t["method_doc"] == "М-27-2023"
    assert t["note"].startswith("распознано 26.1 % (1 строка)")
    assert any(l.get("kind") == "биотест" for l in labs)         # биотест не тронут
    assert any(l.get("_src", "").startswith("другой") for l in labs)  # другой файл
    assert PT.has_table_for(ctx, "13208_26.pdf (листы 1–3)", "", TBO)
    assert PT.has_table_for(ctx, "13208_26.pdf", "", "", protocol_no="13208.26-1-Отх")
    assert not PT.has_table_for(ctx, "другой.pdf", "73310001724", TBO)
    # лучшая запись для отхода
    best = PT.find_for_waste(labs, "7 33 100 01 72 4", TBO)
    assert best is t
    # повторный разбор того же листа — замена, не дубль
    PT.merge_into(ctx, [dict(rec, substances=[{"name": "Бумага", "value": 100.0,
                                                "unit": "%"}], total_pct=100.0,
                             incomplete=False)])
    assert sum(1 for l in ctx.extra["lab_results"] if l.get("method") == "table") == 1


def test_analyzer_store_extras_and_merge_passports_respect_table():
    """ИИ-строки состава того же протокола не дублируют таблицу (rejected с
    причиной); состав паспорта из того же файла берётся из таблицы."""
    from ecodoc.ai.analyzer import ExtractionReport, _merge_passports, _store_extras
    ctx = ReportContext()
    table = {"kind": PT.KIND, "protocol_no": "13208.26-1-Отх", "date": "", "lab": "",
             "lab_attestation": "", "object": TBO, "fkko": "73310001724",
             "method": "table", "method_doc": "",
             "substances": [{"name": "Бумага", "value": 60.0, "unit": "%"},
                            {"name": "Картон", "value": 40.0, "unit": "%"}],
             "_src": "прот.pdf (лист 2)", "page": 2, "total_pct": 100.0,
             "incomplete": False, "note": ""}
    PT.merge_into(ctx, [table])
    rep = ExtractionReport()
    _store_extras(ctx, {"lab_results": [
        {"kind": "КХА", "protocol_no": "13208.26-1-Отх", "date": "17.08.2026",
         "lab": "ИЦ ООО «ТАСИС»", "object": TBO,
         "substances": [{"name": "Картон", "value": "169000", "unit": "мг/кг"}]},
        {"kind": "биотест", "protocol_no": "13308.26-1-ТОтх", "object": TBO}]},
        "прот.pdf (листы 1–3)", rep)
    labs = ctx.extra["lab_results"]
    assert len(labs) == 2 and labs[0]["method"] == "table"
    assert labs[0]["date"] == "17.08.2026" and labs[0]["lab"] == "ИЦ ООО «ТАСИС»"
    assert any("таблицы протокола" in r.reason for r in rep.rejected)
    _merge_passports(ctx, {"waste_passports": [{
        "fkko": "73310001724", "name": TBO, "hazard_class": 4,
        "components": [{"name": "Картон", "percent": "16.9"}]}]},
        "прот.pdf (листы 1–3)", rep)
    p = ctx.extra["waste_passports"][0]
    assert [c["name"] for c in p["components"]] == ["Бумага", "Картон"]
    assert any("из таблицы протокола" in a.field for a in rep.accepted)


def test_passport_prefers_table_and_marks_incomplete(tmp_path):
    """Паспорт: таблица протокола главнее ИИ-строк; неполный состав печатается
    как есть с пометкой «[состав распознан не полностью: …]», без «Прочих»;
    старый программный состав с «Прочими» в waste_details вычищается."""
    from ecodoc.development import waste_passport as wp
    ctx = ReportContext()
    ctx.organization.name = "ООО «Т»"
    ctx.organization.inn = "7800000000"
    ctx.organization.address = "СПб"
    ctx.wastes = [WasteFlow(fkko_code="7 33 100 01 72 4", name=TBO, hazard_class=4)]
    ctx.extra["waste_details"] = {"7 33 100 01 72 4": {
        "components": [{"name": "Прочие компоненты (неидентифицированные)", "percent": "38.60"},
                       {"name": "Картон", "percent": "16.90"}],
        "composition_source": "protocol"}}
    ctx.extra["lab_results"] = [
        {"kind": "КХА", "protocol_no": "13208.26-1-Отх", "date": "17.08.2026",
         "lab": "ИЦ", "object": TBO, "method": "М-27-2023",
         "substances": [{"name": "Картон", "value": "169000", "unit": "мг/кг"}],
         "_src": "п.pdf (листы 1–3)"},
        {"kind": PT.KIND, "protocol_no": "13208.26-1-Отх", "date": "17.08.2026",
         "lab": "ИЦ ООО «ТАСИС»", "lab_attestation": "", "object": TBO,
         "fkko": "73310001724", "method": "table", "method_doc": "М-27-2023",
         "substances": [{"name": "Бумага", "value": 26.1, "unit": "%"},
                        {"name": "Картон", "value": 16.9, "unit": "%"},
                        {"name": "Полиэтилентерефталат", "value": 10.2, "unit": "%"}],
         "_src": "п.pdf (лист 2)", "page": 2, "total_pct": 53.2, "incomplete": True,
         "note": "распознано 53.2 % (3 строки) — проверьте протокол № 13208.26-1-Отх "
                 "(п.pdf, лист 2)"}]
    lab = wp.lab_result_for(ctx, ctx.wastes[0], kinds=wp._COMP_KINDS)
    assert lab["method"] == "table"
    (path,) = wp.generate(ctx, tmp_path)
    d = Document(path)
    text = "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for r in t.rows for c in r.cells])
    assert "Бумага" in text and "26,10" in text and "10,20" in text
    assert "Прочие компоненты" not in text and "38,60" not in text
    assert ("[состав распознан не полностью: распознано 53.2 % (3 строки) — "
            "проверьте протокол № 13208.26-1-Отх (п.pdf, лист 2)]") in text
    assert "М-27-2023" in text and "ТАСИС" in text
    g = wp.gaps(ctx)
    assert any("распознан не полностью" in x and "п.pdf, лист 2" in x for x in g)
    # после генерации в базе — состав без «Прочих», с пометкой
    wp.remember_details(ctx)
    rec = ctx.extra["waste_details"]["7 33 100 01 72 4"]
    assert not any(R.is_synthetic_row(c["name"]) for c in rec["components"])
    assert rec["composition_note"].startswith("[состав распознан не полностью")
    # полный состав — пометки нет, 100 %
    ctx.extra["lab_results"][1]["substances"].append(
        {"name": "Прочее (неклассифицируемые материалы)", "value": 46.8, "unit": "%"})
    ctx.extra["lab_results"][1].update(total_pct=100.0, incomplete=False, note="")
    (path,) = wp.generate(ctx, tmp_path / "b")
    d = Document(path)
    text = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "распознан не полностью" not in text and "46,80" in text
    assert not any("распознан не полностью" in x for x in wp.gaps(ctx))


@pytest.mark.skipif(not Path(r"C:\Users\veter\OneDrive\Формы\Разработка\Протоколы"
                             r"\26212_25,_26312_25т_ООО_ЦЭД_для_ООО_Технострой.pdf").exists(),
                    reason="эталонный скан протокола ЦЭД есть только на машине пользователя")
def test_real_scan_protocol_ced_smoke():
    """Скан ТАСИС для ООО «ЦЭД»/Технострой (12 листов, без текстового слоя):
    три протокола, у каждого свой ФККО; бетон/минвата — 100 %."""
    from ecodoc.parsers.text_extract import _setup_tesseract
    if _setup_tesseract() is None:
        pytest.skip("Tesseract не установлен")
    recs = PT.from_pdf(r"C:\Users\veter\OneDrive\Формы\Разработка\Протоколы"
                       r"\26212_25,_26312_25т_ООО_ЦЭД_для_ООО_Технострой.pdf")
    by_no = {r["protocol_no"]: r for r in recs}
    assert set(by_no) == {"26212.25-1-Отх", "26212.25-2-Отх", "26212.25-3-Отх"}
    assert by_no["26212.25-2-Отх"]["fkko"] == "45711901204"
    assert by_no["26212.25-3-Отх"]["fkko"] == "82240101214"
    assert not by_no["26212.25-2-Отх"]["incomplete"]
    assert not by_no["26212.25-3-Отх"]["incomplete"]
    assert by_no["26212.25-1-Отх"]["fkko"] == "72310101394"
    for r in recs:
        assert not any(R.is_synthetic_row(s["name"]) for s in r["substances"])
