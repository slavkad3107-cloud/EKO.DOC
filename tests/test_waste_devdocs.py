"""Инвентаризация отходов, ПНООЛР, запрос ТУ — настоящие документы по структуре
образцов (Формы/Разработка/{ИНВЕНТАРИЗАЦИЯ ОТХОДОВ, ПНООЛР, ТУ}/ИСТОЧНИК.txt):
разделы, таблицы, отсутствие плейсхолдеров, конкретные gaps."""
from decimal import Decimal

import pytest
from docx import Document

from ecodoc.core.models import NVOSObject, ReportContext, WasteAct, WasteFlow


@pytest.fixture()
def ctx():
    c = ReportContext()
    c.organization.name = "ИП Миних Елена Анатольевна"
    c.organization.short_name = "ИП Миних Е.А."
    c.organization.inn = "780600114472"
    c.organization.ogrn = "307784705100221"
    c.organization.address = "191040, Санкт-Петербург, Коломенская ул., 7"
    c.organization.director_name = "Миних Е.А."
    c.organization.phone = "8 (911) 000-00-00"
    c.period.year = 2025
    c.objects = [NVOSObject(code="41-0247-005048-П", name="База промышленной тары",
                            category="IV", address="Промзона Янино")]
    c.wastes = [
        WasteFlow(fkko_code="47110101521", name="Лампы ртутные", hazard_class=1,
                  generated=Decimal("0.052"), transferred=Decimal("0.052")),
        WasteFlow(fkko_code="73310001724", name="Мусор офисный", hazard_class=4,
                  generated=Decimal("1.9"), transferred=Decimal("1.9")),
        WasteFlow(fkko_code="81111112495", name="Отходы грунта практически неопасные",
                  hazard_class=5, generated=Decimal("120"), transferred=Decimal("120")),
    ]
    c.waste_acts = [
        WasteAct(fkko_code="47110101521", name="Лампы ртутные", mass=Decimal("0.052"),
                 operation="обезвреживание", receiver="ООО «Меркурий»",
                 license="Л020-00113-78/00012345 от 01.02.2020", date="15.02.2025"),
        WasteAct(fkko_code="47110101521", name="Лампы ртутные", mass=Decimal("0.048"),
                 operation="обезвреживание", receiver="ООО «Меркурий»", date="10.03.2024"),
        WasteAct(fkko_code="81111112495", name="Отходы грунта практически неопасные",
                 mass=Decimal("120"), volume_m3=Decimal("75"), operation="размещение",
                 receiver="ООО «Карьер»", date="20.06.2025"),
    ]
    c.extra["waste_passports"] = [{
        "fkko": "47110101521", "name": "Лампы ртутные", "hazard_class": 1,
        "components": [{"name": "стекло", "percent": "92"}]}]
    c.extra["oos_wastes"] = [
        {"stage": "строительство", "fkko": "81111112495",
         "name": "Отходы грунта при проведении открытых земляных работ практически "
                 "неопасные", "hazard_class": 5, "mass_t": "6208", "volume_m3": "3880",
         "density": "1.6", "source_process": "земляные работы",
         "handling": "передача на рекультивацию"},
        {"stage": "эксплуатация", "fkko": "73310001724", "name": "Мусор офисный",
         "hazard_class": 4, "mass_t": "2.1", "volume_m3": "10.5", "density": "0.2"},
    ]
    return c


def _all_text(path) -> str:
    d = Document(path)
    return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        c.text for t in d.tables for r in t.rows for c in r.cells)


def _headings(path) -> list[str]:
    return [p.text for p in Document(path).paragraphs if p.style.name.startswith("Heading")]


# ── инвентаризация отходов ───────────────────────────────────────────────

def test_inventory_docx_structure(ctx, tmp_path):
    from ecodoc.development.waste_inventory import generate_all
    out = generate_all(ctx, tmp_path)
    assert out["path"].endswith(".docx") and len(out["files"]) == 2
    heads = _headings(out["path"])
    for h in ("Введение", "1. Общие сведения об организации и объекте",
              "2. Источники образования отходов",
              "3. Перечень отходов, образующихся на объекте",
              "4. Сводные данные по классам опасности",
              "5. Подтверждающие документы по видам отходов",
              "6. Выводы и рекомендации"):
        assert any(x.startswith(h) for x in heads), h
    text = _all_text(out["path"])
    assert "‹" not in text and "[требуется" not in text      # без плейсхолдеров
    assert "41-0247-005048-П" in text and "ИП Миних Елена Анатольевна" in text
    doc = Document(out["path"])
    # перечень: 3 отхода + шапка; норматив грунта из ООС и факт за 2025 по акту
    table = next(t for t in doc.tables if t.rows[0].cells[1].text.startswith("Наименование отхода"))
    assert len(table.rows) == 4
    soil = next(r for r in table.rows if "8 11 111 12 49 5" in r.cells[2].text)
    assert soil.cells[6].text == "6208" and soil.cells[7].text == "120"
    assert "ООО «Карьер»" in soil.cells[8].text
    # сводная по классам: I, IV, V и ИТОГО
    summary = next(t for t in doc.tables if t.rows[0].cells[0].text == "Класс опасности")
    labels = [r.cells[0].text for r in summary.rows]
    assert labels[-1] == "ИТОГО" and "I класс" in labels and "V класс" in labels
    # лицензия получателя ламп попала в раздел 5
    assert "Л020-00113-78/00012345" in text


def test_inventory_gaps_are_concrete(ctx, tmp_path):
    from ecodoc.development.waste_inventory import gaps
    text = " | ".join(gaps(ctx))
    assert "Мусор офисный: не указан получатель" in text
    assert "нет паспорта отхода" in text
    assert "‹" not in text
    ctx.extra.pop("oos_wastes")
    text = " | ".join(gaps(ctx))
    assert "загрузите раздел ООС или ПНООЛР" in text


# ── ПНООЛР ───────────────────────────────────────────────────────────────

def test_pnoolr_docx_sections_1021(ctx, tmp_path):
    from ecodoc.development.pnoolr import generate_all, rows
    data = {r["fkko"]: r for r in rows(ctx)}
    assert data["81111112495"]["method"] == "project"           # норматив из ООС
    assert data["81111112495"]["norm"] == pytest.approx(6208)
    assert data["47110101521"]["method"] == "fact"              # среднее по актам
    assert data["47110101521"]["norm"] == pytest.approx(0.05)
    out = generate_all(ctx, tmp_path)
    assert out["path"].endswith("ПНООЛР_2025.docx")
    heads = _headings(out["path"])
    for h in ("1. Общие сведения о юридическом лице",
              "2. Сведения о хозяйственной и иной деятельности",
              "3. Сведения об образуемых отходах",
              "4. Расчёт и обоснование нормативов образования отходов",
              "5. Расчёт максимального образования отходов за год",
              "6. Сведения о местах (площадках) накопления отходов",
              "7. Сведения о планируемом обращении с отходами",
              "7.2. Планируемая ежегодная передача отходов другим",
              "7.5. Планируемая ежегодная передача отходов другим хозяйствующим "
              "субъектам с целью их дальнейшего размещения",
              "8. Сводные данные по образованию отходов",
              "Список использованных источников"):
        assert any(x.startswith(h) for x in heads), h
    text = _all_text(out["path"])
    assert "‹" not in text and "[требуется" not in text
    assert "№ 1021" in text
    assert "Но = (0,048 + 0,052) / 2 = 0,05 т/год" in text      # годы по порядку
    # грунт передаётся на размещение — попал в таблицу 7.5 с получателем
    assert "ООО «Карьер»" in text
    gaps_text = " | ".join(out["gaps"])
    assert "объект IV категории — ПНООЛР по ст. 18 ФЗ-89 не требуется" in gaps_text
    assert "пишется экологом" in gaps_text
    assert "не описаны места накопления отходов" in gaps_text


def test_pnoolr_unit_method_reuses_oos_calc(ctx, tmp_path):
    """Исходные данные ООС (численность, материалы, нормативы накопления) →
    расчёт по типовым формулам oos_waste_calc, метод «по удельным нормативам»."""
    from ecodoc.development.pnoolr import calc_rows, generate_docx, rows
    ctx.extra["oos"] = {
        "project": {"workers": 40, "itr": 8, "months": 12},
        "construction": {"materials": [{"name": "Бетон", "kind": "бетон",
                                        "qty": 1000, "unit": "м3"}]},
        "operation": {"wastes_norm": [{"name": "Мусор офисный", "fkko": "73310001724",
                                       "hazard": 4, "count": 10, "count_unit": "чел.",
                                       "norm_m3": 0.95, "density": 0.2}]},
    }
    calc = calc_rows(ctx)
    assert "82220101215" in calc                                # бетон → лом бетона
    assert calc["82220101215"]["t"] == pytest.approx(1000 * 0.3 / 100 * 2.4)
    assert calc["73310001724"]["t"] == pytest.approx(10 * 0.95 * 0.2 + 40 * 0.22 * 0.18
                                                     + 8 * 1.1 * 0.18)
    data = {r["fkko"]: r for r in rows(ctx)}
    assert data["73310001724"]["method"] == "unit"
    out = generate_docx(ctx, tmp_path / "п.docx")
    text = _all_text(out)
    assert "метод расчёта по удельным отраслевым нормативам" in text
    assert "М = 10 чел. × 0,95 м3/год × ρ = 0,2 т/м³" in text
    assert "Справочно — норматив по проектной документации (ООС/ПНООЛР): 2,1 т/год" in text


# ── ТУ ───────────────────────────────────────────────────────────────────

def test_tu_letter_from_oos_construction(ctx, tmp_path):
    from ecodoc.development.tu_waste import gaps, generate, wastes_for_tu
    items = wastes_for_tu(ctx)
    assert [i["fkko"] for i in items] == ["81111112495"]        # только стройка
    assert items[0]["m3"] == 3880 and items[0]["t"] == 6208
    out = generate(ctx, tmp_path / "ту.docx", receiver="ООО «Карьер Рекультивация»",
                   purpose="рекультивации карьера")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    head = " ".join(c.text for r in doc.tables[0].rows for c in r.cells)
    assert "ИП Миних Елена Анатольевна" in text and "ИНН 780600114472" in text
    assert "О выдаче технических условий" in text
    assert "ООО «Карьер Рекультивация»" in head and "Руководителю" in head
    assert "для рекультивации карьера" in text
    cells = [c.text for r in doc.tables[1].rows for c in r.cells]
    assert "8 11 111 12 49 5" in cells and "3880" in cells and "6208" in cells
    assert "передача на рекультивацию" in cells
    assert "Итого: 3880 м³ / 6208 т (по расчёту раздела ООС" in text
    assert "Приложения:" in text and "Индивидуальный предприниматель" in text
    assert "Миних Е.А." in text and "‹" not in text
    g = " | ".join(gaps(ctx, receiver="ООО «Карьер Рекультивация»"))
    assert "адресат" not in g                                  # адресат задан
    assert "нет реквизитов лицензии получателя" in g
    assert "V класс не подтверждён" not in g                   # источник — ООС


def test_tu_letter_falls_back_to_acts(ctx, tmp_path):
    from ecodoc.development.tu_waste import default_receiver, gaps, generate, wastes_for_tu
    ctx.extra.pop("oos_wastes")
    items = wastes_for_tu(ctx)
    assert [i["fkko"] for i in items] == ["81111112495"]        # группа 8 1…
    assert items[0]["t"] == 120 and items[0]["m3"] == pytest.approx(75)  # ρ = 1,6
    assert default_receiver(ctx, items)[0] == "ООО «Карьер»"
    g = " | ".join(gaps(ctx))
    assert "нет нормативов стадии строительства из раздела ООС" in g
    assert "подставлен фактический приёмщик по актам «ООО «Карьер»»" in g
    assert "V класс не подтверждён протоколом биотестирования" in g
    out = generate(ctx, tmp_path / "ту.docx")
    head = " ".join(c.text for r in Document(out).tables[0].rows for c in r.cells)
    assert "ООО «Карьер»" in head
    assert "‹" not in _all_text(out)


def test_tu_gaps_without_anything():
    from ecodoc.development.tu_waste import gaps
    g = " | ".join(gaps(ReportContext()))
    assert "перечень отходов пуст" in g and "не указан адресат" in g
    assert "не указан ИНН" in g and "не задан объект НВОС" in g
