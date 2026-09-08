"""Раунд 7 (08.09.2026): замечание из интерфейса, предпросмотр документов,
статус ООС, плотность из ООС в табличке/сводной, сверка без ложного
замечания «актов нет», проблемы исходников только по базе/каталогу."""
from decimal import Decimal

from ecodoc.core import waste_crosscheck, waste_table, workspace
from ecodoc.core.models import WasteAct, WasteFlow
from ecodoc.gui import preview, server
from ecodoc.intake import candidates, sources


def _site():
    workspace.add_org("ОРГ")
    workspace.add_site("ОРГ", "Пл")
    return workspace.site_dir("ОРГ", "Пл")


def test_feedback_saved_with_context(tmp_path, monkeypatch):
    monkeypatch.setenv("ECODOC_FEEDBACK_DIR", str(tmp_path / "fb"))
    site_dir = _site()
    out = server.api_feedback({}, {"org": "ОРГ", "site": "Пл", "tab": "obj", "subtab": "Отходы",
                                   "text": "не сохраняется удаление", "with_data": True,
                                   "console_errors": ["TypeError: x"], "page_text": "экран"})
    p = tmp_path / "fb"
    md = [x for x in p.glob("*.md") if x.name != "ИНДЕКС.md"]
    assert out["ok"] and len(md) == 1 and out["pending"] == 1
    body = md[0].read_text(encoding="utf-8")
    assert "не сохраняется удаление" in body and "ОЖИДАЕТ" in body and "TypeError" in body
    assert list(p.glob("*_данные.json")) and (p / "ИНДЕКС.md").exists()


def test_doc_preview_docx_and_xlsx(tmp_path):
    import docx
    import openpyxl
    d = docx.Document()
    d.add_paragraph("Паспорт отхода")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Компонент", "%"
    t.cell(1, 0).text, t.cell(1, 1).text = "Бумага <b>", "60"
    f = tmp_path / "п.docx"
    d.save(str(f))
    out = preview.render(f)
    assert out["kind"] == "docx" and "Паспорт отхода" in out["html"]
    assert "&lt;b&gt;" in out["html"] and "<table" in out["html"]
    wb = openpyxl.Workbook()
    wb.active.append(["код", "т"])
    wb.active.append(["73310001724", 1.5])
    x = tmp_path / "т.xlsx"
    wb.save(str(x))
    out2 = preview.render(x)
    assert out2["kind"] == "xlsx" and "73310001724" in out2["html"]
    # API пускает только файлы из папок результатов/базы
    assert "error" in server.api_doc_preview({}, {"path": str(f)})


def test_oos_status_and_density_from_oos():
    site_dir = _site()
    ctx = workspace.load_context("ОРГ", "Пл")
    ctx.wastes = [WasteFlow(fkko_code="81110001495", name="Грунт", hazard_class=5,
                            generated=Decimal("10"))]
    ctx.waste_acts = [WasteAct(fkko_code="81110001495", name="Грунт", mass=Decimal("4"),
                               date="15.03.2025")]
    workspace.save_context("ОРГ", "Пл", ctx)
    sources.remember(site_dir, "f" * 40, file="Раздел ПД № 8 ООС.pdf", method="pdf")
    st = server.api_oos_status({}, {"org": "ОРГ", "site": "Пл"})
    assert st["has_oos_doc"] and not st["has_norms"] and not st["can_reanalyze"]
    assert "был загружен" in st["note"] and "загрузите файл заново" in st["note"]
    ctx.extra["oos_wastes"] = [{"fkko": "81110001495", "name": "Грунт", "stage": "строительство",
                                "mass_t": "10", "volume_m3": "6.25"}]
    workspace.save_context("ОРГ", "Пл", ctx)
    st2 = server.api_oos_status({}, {"org": "ОРГ", "site": "Пл"})
    assert st2["has_norms"] and "1 отход" in st2["note"]
    # плотность из ООС попадает в табличку (акты без объёма)
    dens = waste_table.densities(ctx)
    assert abs(dens["81110001495"] - 1.6) < 1e-6
    rows = {r["fkko"]: r for r in waste_table.rows(ctx)}
    assert rows["81110001495"]["volume"] > 0
    # сверка: «в ООС предусмотрен, актов нет» больше не замечание
    ctx.waste_acts = []
    out = waste_crosscheck.build(ctx, site_dir)
    assert all("справок-актов за период нет" not in i for r in out["rows"] for i in r["issues"])


def test_data_issues_skip_junk_candidate_codes():
    site_dir = _site()
    store = candidates.Store(site_dir)
    # два разных значения для мусорного кода → «сомнение», но код не из ФККО и не в базе
    store.add(candidates.Candidate(key="wastes[fkko=31100000000].generated", value="1", file="a.pdf"))
    store.add(candidates.Candidate(key="wastes[fkko=31100000000].generated", value="2", file="b.pdf"))
    store.save()
    out = server.api_data_issues({}, {"org": "ОРГ", "site": "Пл", "refresh": True})
    assert all("31100000000" not in str(x.get("label", "")) for x in out["categories"]["Отходы"])
