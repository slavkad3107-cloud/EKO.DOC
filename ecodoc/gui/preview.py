"""Предпросмотр сформированного документа прямо в программе.

Эколог попросил видеть результат генерации в интерфейсе, а не только
файлом. Документы всё равно остаются файлами (их сдают), но абзацы и
таблицы .docx / листы .xlsx / текст .xml показываются в панели под кнопкой.
HTML собирается из экранированного текста — никакой разметки из файла.
"""
from __future__ import annotations

import html
from pathlib import Path

MAX_ROWS = 300          # строк на таблицу/лист — дальше «… ещё N»
MAX_CHARS = 400_000


def _esc(s) -> str:
    return html.escape(str(s if s is not None else ""))


def _docx(path: Path) -> str:
    import docx
    from docx.document import Document as _Doc
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    parts: list[str] = []
    body = d.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            tag = "h3" if style.lower().startswith(("heading", "заголовок", "title")) else "p"
            bold = any(r.bold for r in p.runs if r.text.strip())
            cls = ' class="b"' if bold and tag == "p" else ""
            parts.append(f"<{tag}{cls}>{_esc(text)}</{tag}>")
        elif child.tag == qn("w:tbl"):
            t = Table(child, d)
            rows_html = []
            for i, row in enumerate(t.rows):
                if i >= MAX_ROWS:
                    rows_html.append(f"<tr><td colspan=99>… ещё {len(t.rows) - MAX_ROWS} строк</td></tr>")
                    break
                cells = []
                seen = set()
                for c in row.cells:
                    if id(c._tc) in seen:          # объединённые ячейки — один раз
                        continue
                    seen.add(id(c._tc))
                    cells.append(f"<td>{_esc(c.text.strip())}</td>")
                rows_html.append("<tr>" + "".join(cells) + "</tr>")
            parts.append("<table class=\"prev\">" + "".join(rows_html) + "</table>")
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"<h3>Лист «{_esc(ws.title)}»</h3>")
        rows_html = []
        n = 0
        for row in ws.iter_rows(values_only=True):
            if n >= MAX_ROWS:
                rows_html.append("<tr><td colspan=99>… (показаны первые строки)</td></tr>")
                break
            if not any(v not in (None, "") for v in row):
                continue
            n += 1
            rows_html.append("<tr>" + "".join(f"<td>{_esc(v)}</td>" for v in row) + "</tr>")
        parts.append("<table class=\"prev\">" + "".join(rows_html) + "</table>")
    return "\n".join(parts)


def _text(path: Path) -> str:
    data = path.read_text(encoding="utf-8", errors="replace")
    return "<pre>" + _esc(data[:MAX_CHARS]) + ("\n…" if len(data) > MAX_CHARS else "") + "</pre>"


def render(path: str | Path) -> dict:
    """{"html": ..., "kind": "docx|xlsx|xml|txt", "name": ...}"""
    p = Path(path)
    if not p.is_file():
        return {"error": f"файл не найден: {p}"}
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            body, kind = _docx(p), "docx"
        elif ext in (".xlsx", ".xlsm"):
            body, kind = _xlsx(p), "xlsx"
        elif ext in (".xml", ".txt", ".md", ".csv", ".json", ".ics"):
            body, kind = _text(p), ext.lstrip(".")
        else:
            return {"error": f"предпросмотр для {ext} не поддерживается — откройте файл"}
    except Exception as e:                 # битый файл — не валим сервер
        return {"error": f"не удалось прочитать {p.name}: {e}"}
    if len(body) > MAX_CHARS:
        body = body[:MAX_CHARS] + "<p>… (документ обрезан для предпросмотра)</p>"
    return {"html": body, "kind": kind, "name": p.name}
