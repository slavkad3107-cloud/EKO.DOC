"""Детерминированный извлекатель таблиц состава отхода из протоколов
(КХА / морфологический анализ) — без ИИ.

Замечание эколога (08.09.2026): ИИ-разбор текста протокола ИЦ ООО «ТАСИС»
(для ООО «ЦЭД») терял строки таблицы «Результаты исследований» — из 10
компонентов ТБО в базу попали 7 с чужими величинами, а normalize_components
дописывал «Прочие компоненты (неидентифицированные)» до 100 %. В паспорте
такие строки недопустимы: состав надо распознавать полностью, а если не
удалось — явно об этом писать.

Что делает модуль:
  * PDF с текстовым слоем — `page.find_tables()` PyMuPDF (≥ 1.23), таблица
    берётся, если в шапке есть «компонент|наименование|показател|вещество»
    и «содержание|%|масс|мг/кг|г/кг|доля|результат»;
  * PDF-скан / картинка — повторный OCR ТОЛЬКО страниц-кандидатов (в тексте
    страницы есть слова шапки) с координатами слов (`image_to_data`), слова
    собираются в физические строки по вертикали — так строка таблицы
    остаётся строкой, даже если Tesseract разнёс колонки по разным блокам;
    затем строки разбираются регулярным выражением
    «[№] наименование [ед.] число [± погрешность] [методика]»;
  * .docx — таблицы python-docx (контекст — абзацы перед таблицей);
  * .xlsx — листы openpyxl (контекст — ячейки над шапкой);
  * на той же странице (или предыдущей) ищутся наименование отхода
    («Наименование вида отхода: …»), код ФККО (11 цифр с пробелами), номер
    и дата протокола, лаборатория, аттестат, методика.

Результат — записи вида extra.lab_results:
    {kind: "состав/КХА", protocol_no, date, lab, lab_attestation, object,
     fkko, method: "table", method_doc, substances: [{name, value, unit}],
     _src: "файл (лист N)", page: N, total_pct: 61.4, incomplete: bool,
     note: "распознано 61.4 % (7 строк) — проверьте протокол № … лист N"}
Сумма компонентов 99–101 % — состав полный; иначе `incomplete: true` и
никаких «Прочих» не дописывается: паспорт/справка печатают состав как есть с
пометкой, а в gaps попадает конкретная строка с файлом и листом.

`kind` содержит и «состав», и «КХА»: по «состав» его находит отбор протокола
состава (waste_refdata.protocol_matches_waste), по «КХА» — проверка
протоколов IV/V класса (intake/crosscheck.lab_gaps).
"""
from __future__ import annotations

import re
from pathlib import Path

from ecodoc.core.waste_agg import norm_fkko

METHOD = "table"
KIND = "состав/КХА"
COMPLETE_MIN, COMPLETE_MAX = 99.0, 101.0
OCR_MIN_WIDTH = 2000            # px; уже — растягиваем перед OCR

# шапка таблицы состава
_H_NAME = re.compile(r"компонент|наименован|показател|веществ", re.I)
_H_VALUE = re.compile(r"содержан|%|масс|мг/кг|г/кг|дол[яи]\b|результат", re.I)
_H_UNIT = re.compile(r"единиц|ед\.?\s*изм|^\s*ед\.?\s*$", re.I)
_H_NO = re.compile(r"^\s*(№|n|no|п/п|№\s*п/п)\s*(п/п)?\s*$", re.I)
_UNIT_RX = re.compile(r"(мг/кг|г/кг|мас\.?\s*%|%|доли\s*ед\w*|д\.\s*ед\.?)", re.I)
_FKKO_RX = re.compile(r"(?<![\d.,])(\d[  ]?\d{2}[  ]?\d{3}[  ]?\d{2}"
                      r"[  ]?\d{2}[  ]?\d)(?![\d.,])")
_NUM = r"[<>]?\s*\d+(?:[.,]\d+)?"
# строка таблицы в тексте (OCR / текстовый слой без таблицы):
#   [№] Наименование [| % |] 26,100 [± 0,026] [М-27-2023]
_ROW_RX = re.compile(
    r"^\s*[|\[\]]?\s*(?:(?P<idx>\d{1,2})\s*[.)|\]:]?\s+)?[-–—|\[\]]?\s*"
    r"(?P<name>[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё\s(),./\-–—+«»\"']*?)"
    r"\s*[|\[\]:\-–—]*\s*(?P<unit>мг/кг|г/кг|мас\.?\s*%|%)?\s*[|\[\]:\-–—]*\s*"
    r"(?P<val>[<>]?\s*\d+(?:[.,]\d+)?)"
    r"(?:\s*[±+#*]\s*(?P<err>\d+(?:[.,]\d+)?))?"
    r"(?:\s*[|\[\]]*\s*(?P<val2>\d+(?:[.,]\d+)?))?"
    r"(?P<rest>.*)$")
_ZONE_START = re.compile(r"результаты\s+(исслед|измер|компонент)|наименование\s+определ"
                         r"|показател[яи]\s*\(компонент|наименование\s+компонент", re.I)
_ZONE_END = re.compile(r"^\W*(примечани|информаци[яи]\s+о\s+сносках|сотрудник|"
                       r"конец\s+протокола|конец\s+результата|в\s+соответствии\s+с\s+ми)",
                       re.I)
_NOT_ROW = re.compile(r"наименован|показател|результат|погрешност|неопределенност|"
                      r"^\W*(итого|всего|сумма)|примечани|методик|шифр|дата|фкко|"
                      r"аккредит|росс\s*ru|документац|приказ|санитарн|правила|"
                      r"состав\s*\(содержан|морфологическ|химическ\w+\s+состав|"
                      r"^\W*(отклонени|полученные|применяемые|знак|пнд\s*ф|гост|"
                      r"фр\.|ми\s|м-\d)|[«»\"]", re.I)
_OK_REST = re.compile(r"^\W*(протокол|акт|расч[её]т|ми\b|м-\d|пнд|мвИ|гост|фр\.?\s*\d|"
                      r"ру[\s\-]?\d)", re.I)
_MONTHS = {"январ": 1, "феврал": 2, "март": 3, "апрел": 4, "ма": 5, "июн": 6,
           "июл": 7, "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11,
           "декабр": 12}
_PROTOCOL_HEAD = re.compile(r"протокол\s+иссле|протокол\s+испыт|протокол\s+№|"
                            r"протокол\s+измер", re.I)
_LABEL_STOP = re.compile(r"(даты?\s+(осуществ|поступл|отбор)|регистрац|агрегатн|"
                         r"тип\s+пробы|шифр|результат|цель\s+отбор|акт\s+при[её]мки|"
                         r"код\s+(по\s+)?фкко|состав\s+материал|нормативн|"
                         r"физическ\w+\s+форм)", re.I)


# ──────────────────────────── публичный вход ────────────────────────────

def extract(path: str | Path | None, doc=None, ocr: bool = True) -> list[dict]:
    """Записи состава из файла `path` (pdf/docx/xlsx/картинка). `doc` —
    ExtractedDoc того же файла (его постраничный текст нужен для сканов:
    по нему выбираются страницы-кандидаты и берётся контекст; для документа
    из кэша текстов без исходника разбираются только строки текста).
    Ошибки библиотек не рушат приём — возвращается пустой список."""
    p = Path(path) if path else None
    name = p.name if p is not None else Path(getattr(doc, "path", "документ")).name
    try:
        if p is not None and p.is_file():
            suf = p.suffix.lower()
            if suf == ".pdf":
                return from_pdf(p, doc=doc, ocr=ocr)
            if suf == ".docx":
                return from_docx(p)
            if suf in (".xlsx", ".xlsm"):
                return from_xlsx(p)
            if suf in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp") and ocr:
                return from_image(p, name, page=1)
        if doc is not None:
            pages = list(getattr(doc, "pages", None) or [getattr(doc, "text", "") or ""])
            return from_text_pages(pages, name)
    except Exception:
        return []
    return []


def from_pdf(path: str | Path, doc=None, ocr: bool = True) -> list[dict]:
    import fitz

    p = Path(path)
    name = p.name
    doc_pages = list(getattr(doc, "pages", None) or []) if doc is not None else []
    out: list[dict] = []
    prev_text = ""
    with fitz.open(p) as pdf:
        for i, page in enumerate(pdf, 1):
            text = page.get_text("text") or ""
            has_text = len(text.strip()) >= 40
            page_text = text if has_text else (doc_pages[i - 1] if i <= len(doc_pages) else "")
            recs: list[dict] = []
            before = len(out)
            if has_text:
                tables = []
                try:
                    tables = [t.extract() for t in page.find_tables().tables]
                except Exception:
                    tables = []
                for cells in tables:
                    subs, unit_default, hint, unnamed = parse_table(cells)
                    if subs:
                        recs.append(_record(subs, page_text, prev_text, name, i,
                                            unit_default, hint, unnamed))
                if not recs:
                    subs, unit_default, hint, unnamed = parse_rows(page_text.split("\n"))
                    if subs:
                        recs.append(_record(subs, page_text, prev_text, name, i,
                                            unit_default, hint, unnamed))
            elif ocr and (not doc_pages or _candidate(page_text)):
                # скан: слова с координатами → физические строки таблицы
                try:
                    rows_text = _ocr_rows(page.get_pixmap(dpi=300))
                except Exception:
                    rows_text = []
                if rows_text:
                    subs, unit_default, hint, unnamed = parse_rows(rows_text)
                    if subs:
                        ctx_text = "\n".join(rows_text)
                        if page_text.strip():
                            ctx_text = page_text + "\n" + ctx_text
                        recs.append(_record(subs, ctx_text, prev_text, name, i,
                                            unit_default, hint, unnamed))
                        page_text = ctx_text
            _absorb(out, recs, page_text)
            # контекст (номер протокола, отход, ФККО) берётся с предыдущего
            # листа, только если тот — шапка/акт этого же протокола, а не
            # таблица другого отхода
            prev_text = "" if (recs or len(out) != before) else page_text
    return _dedupe(out)


def from_image(path, src_name: str = "", page: int = 1) -> list[dict]:
    """Картинка (jpg/png) или PIL.Image → записи состава через OCR слов."""
    from PIL import Image

    img = path if hasattr(path, "size") else Image.open(path)
    name = src_name or (Path(path).name if not hasattr(path, "size") else "скан")
    rows_text = _ocr_rows_image(img)
    subs, unit_default, hint, unnamed = parse_rows(rows_text)
    if not subs:
        return []
    return [_record(subs, "\n".join(rows_text), "", name, page, unit_default, hint,
                    unnamed)]


def from_docx(path: str | Path) -> list[dict]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    p = Path(path)
    d = Document(str(p))
    out: list[dict] = []
    buf: list[str] = []
    prev = ""
    k = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            buf.append(Paragraph(child, d).text)
        elif tag == "tbl":
            k += 1
            table = Table(child, d)
            cells = [[c.text for c in row.cells] for row in table.rows]
            subs, unit_default, hint, unnamed = parse_table(cells)
            ctx_text = "\n".join(buf)
            if subs:
                out.append(_record(subs, ctx_text, prev, p.name, k, unit_default, hint,
                                   unnamed))
            else:
                # таблица-«шапка» протокола (реквизиты) — тоже контекст
                ctx_text += "\n" + "\n".join("\t".join(r) for r in cells)
            prev = ctx_text
            buf = []
    return _dedupe(out)


def from_xlsx(path: str | Path) -> list[dict]:
    import openpyxl

    p = Path(path)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out: list[dict] = []
    try:
        for k, ws in enumerate(wb.worksheets, 1):
            rows = [["" if c is None else str(c) for c in row]
                    for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(x.strip() for x in r)]
            if not rows:
                continue
            h = _header_index(rows)
            above = "\n".join("\t".join(x for x in r if x.strip()) for r in rows[:h or 0])
            subs, unit_default, hint, unnamed = parse_table(rows)
            if subs:
                out.append(_record(subs, above or "\n".join("\t".join(r) for r in rows),
                                   "", p.name, k, unit_default, hint, unnamed))
    finally:
        wb.close()
    return _dedupe(out)


def from_text_pages(pages: list[str], src_name: str) -> list[dict]:
    """Только текст (кэш текстов, OCR без картинки): строки регэкспом."""
    out: list[dict] = []
    prev = ""
    for i, text in enumerate(pages, 1):
        text = text or ""
        recs = []
        if _candidate(text):
            subs, unit_default, hint, unnamed = parse_rows(text.split("\n"))
            if subs:
                recs.append(_record(subs, text, prev, src_name, i, unit_default, hint,
                                    unnamed))
        _absorb(out, recs, text)
        prev = "" if recs else text
    return _dedupe(out)


# ──────────────────────────── разбор таблиц ────────────────────────────

def _cell(x) -> str:
    return re.sub(r"\s+", " ", norm_text(x)).strip()


def _header_index(rows: list[list]) -> int | None:
    """Индекс строки-шапки (в первых 6 строках; шапка может быть двухэтажной —
    тогда возвращается последняя её строка)."""
    best = None
    for i, row in enumerate(rows[:6]):
        cells = [_cell(c) for c in row]
        joined = " ".join(cells)
        if _H_NAME.search(joined) and (_H_VALUE.search(joined) or _H_UNIT.search(joined)):
            best = i
            # двухэтажная шапка: следующая строка без чисел и тоже «шапочная»
            nxt = rows[i + 1] if i + 1 < len(rows) else None
            if nxt is not None:
                ncells = [_cell(c) for c in nxt]
                if not any(re.search(r"\d", c) for c in ncells if c) and any(
                        _H_VALUE.search(c) or _H_UNIT.search(c) or _H_NAME.search(c)
                        for c in ncells):
                    best = i + 1
            break
    return best


def parse_table(rows: list[list]) -> tuple[list[dict], str, str, int]:
    """Таблица (список строк-ячеек) → (компоненты, единица по умолчанию,
    подсказка о методе, число строк с величиной, но без наименования).
    Пусто — если это не таблица состава."""
    if not rows:
        return [], "", "", 0
    h = _header_index(rows)
    if h is None:
        return [], "", "", 0
    # объединённая (двухэтажная) шапка — текст обоих этажей по колонкам
    ncols = max(len(r) for r in rows)
    head = [""] * ncols
    for r in rows[max(0, h - 1):h + 1]:
        for j, c in enumerate(r):
            head[j] = (head[j] + " " + _cell(c)).strip()
    name_col = next((j for j, c in enumerate(head)
                     if _H_NAME.search(c) and not _H_NO.match(c)), None)
    if name_col is None:
        return [], "", "", 0
    unit_col = next((j for j, c in enumerate(head)
                     if j != name_col and _H_UNIT.search(c)), None)
    value_col = next((j for j, c in enumerate(head)
                      if j not in (name_col, unit_col)
                      and re.search(r"результат|содержан|значени|массов|концентрац", c, re.I)),
                     None)
    if value_col is None:
        value_col = next((j for j, c in enumerate(head)
                          if j not in (name_col, unit_col) and _H_VALUE.search(c)), None)
    if value_col is None:
        # шапка не читается — колонка, где больше всего чисел (не «№ п/п»)
        counts = {}
        for row in rows[h + 1:]:
            for j, c in enumerate(row):
                if j in (name_col, unit_col) or _H_NO.match(head[j] if j < len(head) else ""):
                    continue
                if _parse_num(_cell(c)) is not None:
                    counts[j] = counts.get(j, 0) + 1
        if counts:
            value_col = max(counts, key=lambda j: (counts[j], -j))
    if value_col is None:
        return [], "", "", 0
    m = _UNIT_RX.search(head[value_col])
    unit_default = _norm_unit(m.group(1)) if m else ""
    # вторая колонка величины (например «%» и «мг/кг» рядом)
    value_col2 = next((j for j, c in enumerate(head)
                       if j not in (name_col, unit_col, value_col)
                       and _UNIT_RX.search(c) and not re.search(r"погрешн|неопредел", c, re.I)),
                      None)
    subs: list[dict] = []
    hint = ""
    unnamed = 0
    for row in rows[h + 1:]:
        cells = [_cell(c) for c in row] + [""] * (ncols - len(row))
        filled = [c for c in cells if c]
        if not filled:
            continue
        name = cells[name_col]
        raw = cells[value_col]
        if len(set(filled)) == 1 and not _parse_num(filled[0]):
            # подзаголовок таблицы («Морфологический состав …»)
            hint = hint or filled[0]
            continue
        if not name or _NOT_ROW.search(name):
            if not name and _parse_num(raw) is not None:
                unnamed += 1            # величина есть, наименование потеряно
            continue
        val = _parse_num(raw)
        unit = ""
        if unit_col is not None:
            unit = _norm_unit(cells[unit_col])
        if not unit:
            mu = _UNIT_RX.search(raw)
            unit = _norm_unit(mu.group(1)) if mu else unit_default
        if val is None and value_col2 is not None:
            val = _parse_num(cells[value_col2])
            mu = _UNIT_RX.search(head[value_col2])
            unit = _norm_unit(mu.group(1)) if mu else unit
        if val is None:
            continue
        subs.append({"name": _clean_name(name), "value": val, "unit": unit})
    return subs, unit_default, hint, unnamed


# строка зоны результатов, где OCR потерял наименование: «289000 |Протокол …»,
# «- 28,9 ± 2,0» — величина (с погрешностью/единицей) без слов перед ней
_UNNAMED_RX = re.compile(
    r"^\W{0,3}(?P<val>\d+(?:[.,]\d+)?)\s*(?:[±+#*]\s*\d+(?:[.,]\d+)?|%|мг/кг|г/кг|\|)")


def parse_rows(lines: list[str]) -> tuple[list[dict], str, str, int]:
    """Строки текста (OCR / текстовый слой) → (компоненты, единица по
    умолчанию, подсказка о методе, число строк с величиной без наименования).
    Разбираются только строки в зоне результатов: от шапки таблицы до
    примечаний/подписи."""
    subs: list[dict] = []
    in_zone = False
    unit_default = ""
    hint = ""
    unnamed = 0
    zone_text: list[str] = []
    for raw in _reassemble([norm_text(x) for x in lines]):
        line = raw.strip()
        if not line:
            continue
        if not in_zone:
            if _ZONE_START.search(line):
                in_zone = True
                zone_text.append(line)
            continue
        if _ZONE_END.search(line):
            break
        zone_text.append(line)
        if re.search(r"морфологическ\w+\s+состав|химическ\w+\s+состав", line, re.I) \
                and not hint:
            hint = line
        m = _ROW_RX.match(line)
        if not m:
            um = _UNNAMED_RX.match(line)
            if um and not re.match(r"^\s*(19|20)\d{2}\b", line):
                unnamed += 1
            continue
        name = _clean_name(m.group("name"))
        if (len(name) < 3 or _NOT_ROW.search(name) or len(name.split()) > 8
                or not re.search(r"[а-яёa-z]{3,}", name, re.I)):
            if len(name) < 3 and (m.group("err") or m.group("unit")):
                unnamed += 1
            continue
        rest = (m.group("rest") or "").strip()
        if re.match(r"^[-–]\s*\d", rest):          # «М-27-2023»: число — часть шифра
            continue
        rest_words = re.findall(r"[А-Яа-яЁёA-Za-z]{4,}", rest)
        if rest_words and not _OK_REST.search(rest):
            continue
        val = _parse_num(m.group("val"))
        if val is None:
            continue
        if re.match(r"^\s*(19|20)\d{2}\s*$", m.group("val")) and not m.group("unit"):
            continue                                # год, а не содержание
        unit = _norm_unit(m.group("unit") or "")
        val2 = _parse_num(m.group("val2")) if m.group("val2") else None
        if val2 is not None and not unit:
            # две величины рядом: «100,0 | 1000000» = % и мг/кг
            if abs(val2 - val * 10_000) <= max(1.0, val * 100):
                unit = "%"
            elif abs(val - val2 * 10_000) <= max(1.0, val2 * 100):
                val, unit = val2, "%"
        subs.append({"name": name, "value": val, "unit": unit})
    zone = " ".join(zone_text).lower()
    if "мг/кг" in zone and "%" not in zone:
        unit_default = "мг/кг"
    elif "%" in zone or "масс" in zone or all(s["value"] <= 100 for s in subs):
        unit_default = "%"
    for s in subs:
        if not s["unit"]:
            # без явной единицы: содержание > 100 процентами быть не может —
            # это мг/кг (лист «Результат расчёта…» с потерянной шапкой)
            s["unit"] = "мг/кг" if (unit_default != "мг/кг" and s["value"] > 100) \
                else (unit_default or "%")
    return subs, unit_default, hint, unnamed


_IDX_LINE = re.compile(r"^\s*\d{1,2}\s*[.)]?\s*$")
_VAL_LINE = re.compile(rf"^\s*{_NUM}\s*(?:[±+#*]\s*\d+(?:[.,]\d+)?)?\s*$")


def _reassemble(lines: list[str]) -> list[str]:
    """Текстовый слой PDF без линий таблицы: ячейки строки идут отдельными
    строками («1» / «Бумага» / «%» / «26,100 ± 0,026») — собрать их в одну."""
    out: list[str] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        if _IDX_LINE.match(line) and i + 1 < n and not _VAL_LINE.match(lines[i + 1]):
            group = [line.strip()]
            j = i + 1
            while j < n and j - i <= 4:
                group.append(lines[j].strip())
                if _VAL_LINE.match(lines[j]):
                    break
                j += 1
            if _VAL_LINE.match(group[-1]) and len(group) >= 3:
                out.append(" | ".join(group))
                i = j + 1
                continue
        out.append(line)
        i += 1
    return out


def _candidate(text: str) -> bool:
    return bool(text and _H_NAME.search(text) and _H_VALUE.search(text)
                and _ZONE_START.search(text))


def _parse_num(s) -> float | None:
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).strip().replace(" ", "").replace(" ", "")
    if not t:
        return None
    if t.startswith("<") or t.lower().startswith("менее"):
        return 0.0
    m = re.match(r"^[<>]?(\d+(?:[.,]\d+)?)", t)
    if not m:
        return None
    try:
        return float(m.group(1).replace(",", "."))
    except ValueError:
        return None


def _norm_unit(u: str) -> str:
    u = (u or "").lower().replace(" ", "")
    if "мг/кг" in u:
        return "мг/кг"
    if "г/кг" in u:
        return "г/кг"
    if "%" in u:
        return "%"
    if "дол" in u or "д.ед" in u:
        return "доли"
    return ""


def _clean_name(s: str) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip(" -–—|[]:;,.")
    return s


# ──────────────────────────── контекст протокола ────────────────────────────

def norm_text(s: str) -> str:
    """Мягкие/неразрывные дефисы и пробелы из текстового слоя PDF → обычные."""
    return (str(s or "").replace("\xad", "-").replace("‑", "-").replace("‐", "-")
            .replace("−", "-").replace("\xa0", " "))


def context_info(text: str, prev_text: str = "") -> dict:
    """Реквизиты протокола и отхода из текста страницы (и предыдущей)."""
    info = {"protocol_no": "", "date": "", "lab": "", "lab_attestation": "",
            "object": "", "fkko": "", "method_doc": ""}
    for src, primary in ((norm_text(text), True), (norm_text(prev_text), False)):
        if not src:
            continue
        if not info["protocol_no"]:
            no, date = _protocol_no(src)
            if no:
                info["protocol_no"], info["date"] = no, date
        if not info["object"]:
            obj = _waste_name(src)
            if obj and (primary or len(obj) > 8):
                info["object"] = obj
        if not info["fkko"]:
            m = _FKKO_RX.search(src)
            if m:
                info["fkko"] = norm_fkko(m.group(1))
        if not info["lab"]:
            info["lab"] = _lab(src)
        if not info["lab_attestation"]:
            m = re.search(r"(РОСС\s*RU[.\s]*\d{4}[.\s]*[\dА-ЯA-Zа-яa-z]{2,8})", src)
            if m:
                info["lab_attestation"] = "№ " + re.sub(r"\s+", " ", m.group(1)).strip()
        if not info["method_doc"]:
            m = re.search(r"(М-\d+-\d{4}|ПНД\s*Ф\s*[\d.:\-]+|ФР\.\s*1\.\d+\.\d+\.\d+)", src)
            if m:
                info["method_doc"] = re.sub(r"\s+", " ", m.group(1))
        if primary and not info["object"]:
            continue
    return info


def _protocol_no(text: str) -> tuple[str, str]:
    m = _PROTOCOL_HEAD.search(text)
    windows = []
    if m:
        windows.append(text[m.start():m.start() + 400])
    # «протокол исследований (измерений) № 26212.25-1-Отх от 19 декабря 2025 г.»
    windows.append(text)
    for win in windows:
        mm = re.search(r"№\s*(?P<no>\d(?:[\d .]*\d)?(?:\s*[-–]\s*\d+)*"
                       r"(?:\s*[-–]\s*[А-Яа-яA-Za-z]{1,6})?)"
                       r"\s*(?:от\s+(?P<date>\d{1,2}\s+[а-яё]+\s+\d{4}|\d{2}\.\d{2}\.\d{4}))?",
                       win)
        if mm and (m or "протокол" in win[:mm.start()].lower()[-120:]):
            no = mm.group("no")
            tail = win[mm.end("no"):]
            # номер, разорванный переносом строки OCR: «№ 26212» / «.25 -1 -Отх»
            cont = re.match(r"\s*\n\s*([.\-–]\s*\d[\d .\-–]*(?:[А-Яа-яA-Za-z]{1,6})?)", tail)
            if cont:
                no += cont.group(1)
            no = norm_protocol_no(no)
            # «№ 1» посреди текста — не номер; «№ 5-Отх», «№ 20002.25-1-Отх» — номер
            if re.search(r"\d{2,}", no) or re.search(r"\d-\S", no):
                return no, _norm_date(mm.group("date") or "")
    return "", ""


def norm_protocol_no(s: str) -> str:
    """«13208 .26 -1 -Отх» → «13208.26-1-Отх»."""
    s = re.sub(r"\s+", "", str(s or ""))
    return s.replace("–", "-").strip("-.,")


def _norm_date(s: str) -> str:
    s = (s or "").strip()
    if not s:
        return ""
    if re.match(r"^\d{2}\.\d{2}\.\d{4}$", s):
        return s
    m = re.match(r"^(\d{1,2})\s+([а-яё]+)\s+(\d{4})$", s.lower())
    if not m:
        return s
    mon = next((n for k, n in _MONTHS.items() if m.group(2).startswith(k)), 0)
    if not mon:
        return s
    return f"{int(m.group(1)):02d}.{mon:02d}.{m.group(3)}"


def _junk(line: str) -> bool:
    """Строка OCR-мусора: букв кириллицы меньше 60 % значимых символов."""
    s = re.sub(r"\s+", "", line)
    if not s:
        return True
    cyr = len(re.findall(r"[а-яёА-ЯЁ]", s))
    return cyr / len(s) < 0.6


def _waste_name(text: str) -> str:
    # «вила»/«отхола» — типичные ошибки OCR («д» → «л»)
    for rx in (r"наименование\s+в[ил]\w{1,2}\s+(?:отх\w{2,4}\s*\*?\s*[:;]?\s*)?"
               r"(?P<n>[А-ЯЁ][^\n]*(?:\n[^\n]+){0,2})",
               r"код\s+по\s+фкко[^\n]*\n(?P<n>[^\n]+)",
               r"(?:^|\n)\s*отход\w*\s*:\s*(?P<n>[А-Яа-яЁё][^\n]{7,})"):
        m = re.search(rx, text, re.I)
        if not m:
            continue
        lines = m.group("n").split("\n")
        keep = [lines[0]]
        for extra in lines[1:]:
            # продолжение наименования: не подпись поля, не мусор OCR,
            # начинается со строчной буквы или скобки
            if (_LABEL_STOP.search(extra) or _junk(extra)
                    or not re.match(r"^\s*[а-яёa-z(]", extra) or ":" in extra):
                break
            keep.append(extra)
        n = " ".join(keep)
        cut = _LABEL_STOP.search(n)
        if cut:
            n = n[:cut.start()]
        n = re.sub(r"\s+", " ", n).strip(" :;*—-|")
        n = re.sub(r"^(отход\w*\s*:\s*)", "", n, flags=re.I)
        if len(n) >= 8 and re.search(r"[а-яё]{3,}", n, re.I) and not _junk(n):
            return n[:200]
    return ""


def _lab(text: str) -> str:
    for rx in (r"((?:ИЦ|ИЛ|ИЛЦ)\s+(?:ООО|АО|ЗАО|ФБУЗ|ФГБУ)\s*[«\"“][^»\"”\n]{2,60}[»\"”])",
               r"(Испытательн\w+\s+(?:центр|лаборатори\w+)\s+(?:ООО|АО|ЗАО|ФБУЗ|ФГБУ)\s*"
               r"[«\"“][^»\"”\n]{2,60}[»\"”])"):
        m = re.search(rx, text, re.I)
        if m:
            s = re.sub(r"\s+", " ", m.group(1)).strip()
            return s.replace('"', "«", 1).replace('"', "»", 1) if s.count('"') == 2 else s
    return ""


# ──────────────────────────── записи и полнота ────────────────────────────

def total_percent(substances: list[dict]) -> float:
    """Сумма компонентов в % (мг/кг ÷ 10 000, г/кг ÷ 10, доли × 100);
    без единицы — по величине суммы (> 150 → мг/кг), как в
    waste_refdata.detect_unit."""
    from ecodoc.core import waste_refdata as R

    subs = [s for s in substances if isinstance(s, dict)]
    unit_fallback = R.detect_unit([{"unit": "", "percent": s.get("value")}
                                   for s in subs if not s.get("unit")]) if subs else "%"
    total = 0.0
    for s in subs:
        v = R.parse_value(s.get("value"))
        if v is None:
            continue
        unit = _norm_unit(str(s.get("unit") or "")) or unit_fallback
        total += v / R._DIV.get(unit, 1.0)
    return round(total, 2)


def is_complete(total: float) -> bool:
    return COMPLETE_MIN <= total <= COMPLETE_MAX


def _record(subs: list[dict], text: str, prev_text: str, src_name: str, page: int,
            unit_default: str, hint: str, unnamed: int = 0) -> dict:
    info = context_info(text, prev_text)
    total = total_percent(subs)
    rec = {"kind": KIND, "protocol_no": info["protocol_no"], "date": info["date"],
           "lab": info["lab"], "lab_attestation": info["lab_attestation"],
           "object": info["object"], "fkko": info["fkko"], "method": METHOD,
           "method_doc": info["method_doc"], "substances": subs,
           "unit": unit_default, "_src": f"{src_name} (лист {page})", "page": page,
           "total_pct": total, "incomplete": not is_complete(total), "note": ""}
    if hint:
        rec["analysis"] = hint
    if unnamed:
        rec["rows_unnamed"] = int(unnamed)
    if rec["incomplete"]:
        rec["note"] = incomplete_note(rec)
    return rec


def _plural(n: int, one: str, few: str, many: str) -> str:
    n = abs(int(n))
    if n % 10 == 1 and n % 100 != 11:
        return f"{n} {one}"
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return f"{n} {few}"
    return f"{n} {many}"


def incomplete_note(rec: dict) -> str:
    """«распознано 61.4 % (7 строк) — проверьте протокол № … (файл, лист N)»;
    если часть строк прочитана без наименования (OCR) — и об этом."""
    n = len(rec.get("substances") or [])
    no = rec.get("protocol_no") or "без номера"
    rows = _plural(n, "строка", "строки", "строк")
    unnamed = int(rec.get("rows_unnamed") or 0)
    if unnamed:
        rows += (f", ещё {_plural(unnamed, 'строка', 'строки', 'строк')} "
                 f"с величиной без наименования")
    return (f"распознано {rec.get('total_pct', 0):.1f} % ({rows}) — проверьте "
            f"протокол № {no} ({src_file(rec.get('_src'))}, лист {rec.get('page') or '?'})")


def _absorb(out: list[dict], recs: list[dict], page_text: str) -> None:
    """Продолжение таблицы на следующей странице (нет своей шапки протокола
    и наименования отхода) — дописать к предыдущей записи."""
    for rec in recs:
        prev = out[-1] if out else None
        if (prev is not None and not rec.get("protocol_no") and not rec.get("object")
                and not _PROTOCOL_HEAD.search(page_text or "")
                and prev.get("page") == rec["page"] - 1):
            prev["substances"] = list(prev["substances"]) + list(rec["substances"])
            prev["total_pct"] = total_percent(prev["substances"])
            prev["incomplete"] = not is_complete(prev["total_pct"])
            prev["_src"] = f"{src_file(prev['_src'])} (листы {prev['page']}–{rec['page']})"
            prev["note"] = incomplete_note(prev) if prev["incomplete"] else ""
            continue
        out.append(rec)


def _dedupe(recs: list[dict]) -> list[dict]:
    """Одна запись на протокол: если таблица одного и того же протокола
    встретилась дважды (лист протокола и лист «Результат расчёта»), остаётся
    более полная, реквизиты (ФККО/наименование) дополняются из второй."""
    out: list[dict] = []
    for rec in recs:
        dup = next((x for x in out if same_protocol(x, rec)), None)
        if dup is None:
            out.append(rec)
            continue
        keep, drop = (dup, rec) if _better(dup) >= _better(rec) else (rec, dup)
        for k in ("fkko", "object", "protocol_no", "date", "lab", "lab_attestation",
                  "method_doc"):
            if not keep.get(k) and drop.get(k):
                keep[k] = drop[k]
        keep["note"] = incomplete_note(keep) if keep.get("incomplete") else ""
        if keep is rec:
            out[out.index(dup)] = rec
    return out


def _better(a: dict) -> tuple:
    """Оценка записи: полная > неполная, больше распознанных процентов >
    меньше, больше строк > меньше (реквизиты — ФККО, номер — всё равно
    дополняются из отброшенной записи)."""
    return (0 if a.get("incomplete") else 1,
            min(float(a.get("total_pct") or 0.0), 100.0),
            len(a.get("substances") or []))


# ──────────────────────────── сопоставление ────────────────────────────

_STOP = {"отходы", "отход", "прочие", "прочая", "прочий", "изделия", "изделий",
         "или", "для", "при", "менее", "более", "виде", "иных", "иные",
         "незагрязненные", "незагрязненный", "незагрязненная", "утратившие",
         "утративший", "утратившая", "потребительские", "свойства", "форме",
         "количестве"}


def norm_name(s: str) -> str:
    """Нижний регистр, ё→е, без скобок и пунктуации, один пробел."""
    s = str(s or "").lower().replace("ё", "е")
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"[^a-zа-я0-9\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _words(s: str) -> set[str]:
    return {w for w in norm_name(s).split() if len(w) > 3 and w not in _STOP}


def names_match(a: str, b: str) -> bool:
    """Одно и то же наименование отхода: первые 40 символов нормализованных
    строк совпадают либо пересечение слов длиннее 3 букв ≥ 70 %."""
    na, nb = norm_name(a), norm_name(b)
    if not na or not nb or len(na) < 8 or len(nb) < 8:
        return False
    if na[:40] == nb[:40]:
        return True
    wa, wb = _words(na), _words(nb)
    if not wa or not wb:
        return False
    return len(wa & wb) / min(len(wa), len(wb)) >= 0.7


def match_waste(rec: dict, fkko, name: str = "") -> bool:
    """Запись протокола — про этот отход: по коду ФККО, иначе по наименованию."""
    if not isinstance(rec, dict):
        return False
    code = norm_fkko(fkko or "")
    rcode = norm_fkko(rec.get("fkko") or "")
    if code and rcode:
        return code == rcode
    target = " ".join(str(rec.get(k) or "") for k in ("object", "waste", "name"))
    if code and len(code) == 11 and code in re.sub(r"\D", "", target):
        return True
    return names_match(target, name)


def same_protocol(a: dict, b: dict) -> bool:
    """Та же таблица того же документа: номер протокола, иначе ФККО/имя."""
    if src_file(a.get("_src")) != src_file(b.get("_src")):
        return False
    na, nb = norm_protocol_no(a.get("protocol_no")), norm_protocol_no(b.get("protocol_no"))
    if na and nb:
        return na.lower() == nb.lower()
    if na or nb:
        return False
    return match_waste(a, b.get("fkko"), b.get("object") or "")


def src_file(label) -> str:
    return str(label or "").split(" (лист")[0].strip()


def is_composition_kind(kind) -> bool:
    k = str(kind or "").lower()
    return any(x in k for x in ("кха", "хим", "морф", "состав", "компонент"))


def find_for_waste(labs: list, fkko, name: str = "") -> dict | None:
    """Лучшая табличная запись состава для отхода: полная > неполная,
    больше компонентов > меньше."""
    cands = [l for l in labs or [] if isinstance(l, dict) and l.get("method") == METHOD
             and l.get("substances") and match_waste(l, fkko, name)]
    if not cands:
        return None
    return sorted(cands, key=lambda l: (bool(l.get("incomplete")),
                                        -len(l.get("substances") or [])))[0]


# ──────────────────────────── слияние в контекст ────────────────────────────

def merge_into(ctx, records: list[dict]) -> dict:
    """Записать табличные записи в extra.lab_results. Табличный результат
    главнее ИИ-строк того же документа: ИИ-записи состава с тем же номером
    протокола / отходом удаляются (их дата, лаборатория, аттестат
    дописываются в табличную запись, если у неё этого нет); повторный разбор
    того же листа заменяет прежнюю табличную запись."""
    if not records:
        return {"added": 0, "replaced": 0}
    if not isinstance(ctx.extra, dict):
        ctx.extra = {}
    labs = ctx.extra.setdefault("lab_results", [])
    added = replaced = 0
    for rec in records:
        keep: list = []
        for old in labs:
            if not (isinstance(old, dict) and is_composition_kind(old.get("kind"))
                    and same_protocol(old, rec)):
                keep.append(old)
                continue
            replaced += 1
            for k in ("date", "lab", "lab_attestation", "object", "fkko"):
                if not rec.get(k) and old.get(k):
                    rec[k] = old[k]
            # имя лаборатории из OCR с латиницей («TACHC») — взять из ИИ-записи
            if old.get("lab") and re.search(r"[A-Za-z]", str(rec.get("lab") or "")):
                rec["lab"] = old["lab"]
            if old.get("method") != METHOD and not rec.get("method_doc") and old.get("method"):
                rec["method_doc"] = str(old["method"])
        labs[:] = keep
        if rec.get("incomplete"):
            rec["note"] = incomplete_note(rec)
        labs.append(rec)
        added += 1
    return {"added": added, "replaced": replaced}


def find_in_file(ctx, src: str, fkko="", name: str = "",
                 protocol_no: str = "") -> dict | None:
    """Табличная запись из этого документа про этот отход (по номеру
    протокола, иначе по ФККО/наименованию) — для analyzer: состав ИИ из
    того же файла не берём, таблица главнее."""
    extra = ctx.extra if isinstance(getattr(ctx, "extra", None), dict) else {}
    file = src_file(src)
    pno = norm_protocol_no(protocol_no).lower()
    for l in extra.get("lab_results") or []:
        if not (isinstance(l, dict) and l.get("method") == METHOD
                and src_file(l.get("_src")) == file):
            continue
        lno = norm_protocol_no(l.get("protocol_no")).lower()
        if pno and lno:
            if pno == lno:
                return l
            continue
        if match_waste(l, fkko, name):
            return l
    return None


def has_table_for(ctx, src: str, fkko="", name: str = "", protocol_no: str = "") -> bool:
    return find_in_file(ctx, src, fkko, name, protocol_no) is not None


def as_components(rec: dict) -> list[dict]:
    """Компоненты табличной записи в виде состава паспорта
    ({name, percent, unit}) — как их ждут waste_passports/waste_details."""
    return [{"name": str(s.get("name") or ""), "percent": s.get("value", ""),
             "unit": str(s.get("unit") or "")}
            for s in (rec.get("substances") or []) if isinstance(s, dict)
            and str(s.get("name") or "").strip()]


# ──────────────────────────── OCR строк ────────────────────────────

def _ocr_rows(pix) -> list[str]:
    from io import BytesIO

    from PIL import Image

    return _ocr_rows_image(Image.open(BytesIO(pix.tobytes("png"))))


def _ocr_rows_image(img) -> list[str]:
    """Слова с координатами → физические строки (по вертикали), слова в
    строке — слева направо. Строки таблицы остаются строками, даже если
    Tesseract разнёс колонки по разным блокам."""
    from ecodoc.parsers.text_extract import _setup_tesseract

    lang = _setup_tesseract()
    if lang is None:
        return []
    import pytesseract

    # снимок листа низкого разрешения (~100 dpi, как в pages/) Tesseract
    # читает плохо — растянуть до ширины ~300 dpi A4
    try:
        if img.width and img.width < OCR_MIN_WIDTH:
            from PIL import Image

            k = OCR_MIN_WIDTH / float(img.width)
            img = img.resize((int(img.width * k), int(img.height * k)), Image.LANCZOS)
    except Exception:
        pass
    try:
        data = pytesseract.image_to_data(img, lang=lang, timeout=300,
                                         output_type=pytesseract.Output.DICT)
    except Exception:
        return []
    # 1) слова → строки Tesseract (внутри блока он сам учитывает наклон)
    segs: dict[tuple, list] = {}
    for j in range(len(data.get("text") or [])):
        t = (data["text"][j] or "").strip()
        try:
            conf = float(data["conf"][j])
        except (TypeError, ValueError):
            conf = -1
        if not t or conf < 0:
            continue
        key = (data["block_num"][j], data["par_num"][j], data["line_num"][j])
        x0, y0 = int(data["left"][j]), int(data["top"][j])
        x1, y1 = x0 + int(data["width"][j]), y0 + int(data["height"][j])
        seg = segs.setdefault(key, [x0, y0, x1, y1, []])
        seg[0], seg[1] = min(seg[0], x0), min(seg[1], y0)
        seg[2], seg[3] = max(seg[2], x1), max(seg[3], y1)
        seg[4].append((x0, t))
    if not segs:
        return []
    items = sorted(segs.values(), key=lambda s: ((s[1] + s[3]) / 2.0, s[0]))
    # 2) отрезки разных колонок на одной физической строке: вертикальные
    #    интервалы перекрываются ≥ 50 % меньшей высоты, по горизонтали не
    #    пересекаются — это одна строка таблицы
    rows: list[list] = []          # [y0, y1, [(x0, x1, words)]]
    for x0, y0, x1, y1, words in items:
        h = max(1, y1 - y0)
        target = None
        for row in rows:
            ov = min(row[1], y1) - max(row[0], y0)
            if ov < 0.5 * min(h, max(1, row[1] - row[0])):
                continue
            if any(min(sx1, x1) - max(sx0, x0) > 0.2 * min(x1 - x0, sx1 - sx0)
                   for sx0, sx1, _ in row[2]):
                continue
            target = row
            break
        if target is None:
            rows.append([y0, y1, [(x0, x1, words)]])
        else:
            target[0], target[1] = min(target[0], y0), max(target[1], y1)
            target[2].append((x0, x1, words))
    rows.sort(key=lambda r: (r[0] + r[1]) / 2.0)
    out = []
    for _y0, _y1, segments in rows:
        segments.sort()
        out.append(" ".join(" ".join(t for _, t in sorted(ws)) for _, _, ws in segments))
    return out
